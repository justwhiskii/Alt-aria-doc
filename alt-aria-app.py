import streamlit as st
import google.generativeai as genai
from PIL import Image
import json
import io
import re

# Import docx safely to prevent top-level app crashes
try:
    from docx import Document
    from docx.shared import Inches
except ImportError:
    st.error("Missing dependency: `python-docx`. Please add `python-docx` to your requirements.txt file.")
    st.stop()

# --- PAGE CONFIGURATION ---
st.set_page_config(page_title="Email A11y Generator", layout="wide", page_icon="♿")

st.title("♿ Email Accessibility (A11y) Generator")
st.caption("Generate Deque-compliant Alt Text, ARIA Labels, and export directly to Word/Google Docs.")

# --- FETCH API KEY FROM STREAMLIT SECRETS ---
if "GEMINI_API_KEY" in st.secrets:
    api_key = st.secrets["GEMINI_API_KEY"]
else:
    st.error("⚠️ API Key not found! Please add `GEMINI_API_KEY` to Streamlit App Settings -> Secrets.")
    st.stop()

# --- HELPER: CLEAN GEMINI JSON OUTPUT ---
def parse_gemini_json(text):
    """Strips Markdown formatting like ```json ... ``` before parsing."""
    cleaned = re.sub(r'```(?:json)?\n?', '', text).strip()
    cleaned = cleaned.rstrip('`')
    return json.loads(cleaned)

# --- DYNAMIC MODEL GENERATOR ---
def generate_a11y_data(api_key, prompt, image):
    genai.configure(api_key=api_key)
    
    candidate_models = []
    
    # Query API for active models available to this key
    try:
        for m in genai.list_models():
            if 'generateContent' in m.supported_generation_methods:
                model_name = m.name.replace('models/', '')
                if 'flash' in model_name or 'pro' in model_name:
                    candidate_models.append(model_name)
    except Exception:
        pass

    defaults = ['gemini-1.5-flash', 'gemini-1.5-pro']
    for d in defaults:
        if d not in candidate_models:
            candidate_models.append(d)

    last_error = None
    for model_name in candidate_models:
        try:
            model = genai.GenerativeModel(
                model_name=model_name,
                generation_config={"response_mime_type": "application/json"}
            )
            response = model.generate_content([prompt, image])
            return response.text, model_name
        except Exception as e:
            last_error = e
            continue

    raise Exception(f"All model attempts failed. Last error: {str(last_error)}")

# --- HELPER FUNCTION: CREATE WORD DOC ---
def create_word_doc(a11y_data, original_image):
    doc = Document()
    doc.add_heading('Accessibility A11y Documentation', 0)
    
    table = doc.add_table(rows=1, cols=5)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    headers = ['Section of Email', 'Block Description', 'Visual', 'Alt Text', 'ARIA Label']
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        
    img_width, img_height = original_image.size
    
    for row_data in a11y_data:
        row_cells = table.add_row().cells
        row_cells[0].text = str(row_data.get("section", "—"))
        row_cells[1].text = str(row_data.get("description", "—"))
        
        box = row_data.get("box_2d")
        if box and isinstance(box, list) and len(box) == 4:
            ymin, xmin, ymax, xmax = box
            left = (xmin / 1000.0) * img_width
            top = (ymin / 1000.0) * img_height
            right = (xmax / 1000.0) * img_width
            bottom = (ymax / 1000.0) * img_height
            
            cropped_slice = original_image.crop((left, top, right, bottom))
            
            img_stream = io.BytesIO()
            cropped_slice.save(img_stream, format='PNG')
            img_stream.seek(0)
            
            paragraph = row_cells[2].paragraphs[0]
            run = paragraph.add_run()
            run.add_picture(img_stream, width=Inches(1.2))
        else:
            row_cells[2].text = "—"
            
        row_cells[3].text = str(row_data.get("alt_text", "n/a"))
        row_cells[4].text = str(row_data.get("aria_label", "n/a"))
        
    doc_stream = io.BytesIO()
    doc.save(doc_stream)
    doc_stream.seek(0)
    return doc_stream

# --- INPUT SECTION ---
col_left, col_right = st.columns([1, 1])

with col_left:
    uploaded_file = st.file_uploader("1. Upload Email Proof (PNG, JPG, JPEG)", type=["png", "jpg", "jpeg"])

with col_right:
    developer_notes = st.text_area(
        "2. Custom Callouts / Build Context", 
        height=150,
        placeholder="e.g.,\n- Hero is an animated GIF\n- Module 2 is a Movable Ink block"
    )

if uploaded_file:
    image = Image.open(uploaded_file)
    st.image(image, caption="Uploaded Email Proof", width=300)
    
    if st.button("🚀 Generate A11y Table", type="primary"):
        prompt = f"""
        You are an expert Accessibility (A11y) Specialist. Analyze this email proof and generate data for a strict Deque-compliant A11y table.

        DEVELOPER CALLOUTS: {developer_notes if developer_notes else "None provided."}

        RULES:
        1. Decorative Lifestyle Photos: SKIP ROW completely unless it contains incidental text or an imposed brand logo.
        2. Essential Graphics & Logos: Provide clear, descriptive Alt Text.
        3. ARIA Labels: Must literally match the exact text in the button or link.
        4. Unclickable Graphics: ARIA Label must be "n/a - not clickable".
        5. Movable Ink: Group into a single row, set Block Description to "Creative Optimizer", generalized Alt/ARIA text.
        6. Footers: Group entire text links into a single row.
        7. Visual Bounding Boxes: Provide normalized coordinates [ymin, xmin, ymax, xmax] (0 to 1000 scale).

        Return ONLY a JSON array containing objects with these exact keys:
        [
          {{"section": "...", "description": "...", "box_2d": [ymin, xmin, ymax, xmax], "alt_text": "...", "aria_label": "..."}}
        ]
        """

        with st.spinner("Analyzing email structure, extracting ARIA labels, and slicing image regions..."):
            try:
                raw_response, used_model = generate_a11y_data(api_key, prompt, image)
                a11y_data = parse_gemini_json(raw_response)
                
                st.success(f"Analysis Complete (Model: {used_model})!")
                
                word_file = create_word_doc(a11y_data, image)
                
                st.download_button(
                    label="📄 Download Word Document (Upload to Google Drive)",
                    data=word_file,
                    file_name="A11y_Documentation.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    type="primary"
                )
                
            except Exception as e:
                st.error(f"Error during execution: {str(e)}")
